"""Render a ReportDocument (web/lib/report/document.ts, format 1) to Word or PDF.

Figures are drawn with matplotlib (shipped with Starsim). The PDF uses
Helvetica with the exporter's Latin fallback, or reportlab's built-in CJK font
when the text needs it (the helpers in epichat.exporter).
"""
from __future__ import annotations

import io
import json
from typing import Any

from epichat.exporter import _has_cjk, _rl_register_cjk, _sanitize_latin

MEDIA = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
MAX_BYTES = 2_000_000
MAX_SECTIONS = 40
MAX_LINES = 12
BLOCK_KINDS = {"paragraph", "bullets", "table", "keyValues", "figure", "note"}
COLORS = ["#8c3b2a", "#2b5f8a", "#3b7d4f", "#b8860b", "#6a4c93", "#c0392b",
          "#16a085", "#7f8c8d", "#d35400", "#2c3e50", "#8e44ad", "#27ae60"]


def validate_document(doc: Any) -> list[str]:
    """Everything the renderers rely on; the route answers 422 with these."""
    if not isinstance(doc, dict):
        return ["document must be an object"]
    problems: list[str] = []
    if doc.get("version") != 1:
        problems.append("version must be 1")
    for key in ("title", "subtitle", "generatedAt", "language"):
        if not isinstance(doc.get(key), str):
            problems.append(f"{key} must be a string")
    sections = doc.get("sections")
    if not isinstance(sections, list):
        return problems + ["sections must be a list"]
    if len(sections) > MAX_SECTIONS:
        problems.append(f"at most {MAX_SECTIONS} sections")
    for s_i, section in enumerate(sections):
        if not isinstance(section, dict) or not isinstance(section.get("heading"), str) or not isinstance(section.get("blocks"), list):
            problems.append(f"section {s_i} needs a heading and blocks")
            continue
        for b_i, block in enumerate(section["blocks"]):
            where = f"section {s_i} block {b_i}"
            kind = block.get("kind") if isinstance(block, dict) else None
            if kind not in BLOCK_KINDS:
                problems.append(f"{where}: unknown kind {kind!r}")
                continue
            if kind in ("paragraph", "note") and not isinstance(block.get("text"), str):
                problems.append(f"{where}: text must be a string")
            if kind == "bullets" and not all(isinstance(i, str) for i in block.get("items", [])):
                problems.append(f"{where}: items must be strings")
            if kind == "keyValues" and not all(isinstance(i, list) and len(i) == 2 for i in block.get("items", [])):
                problems.append(f"{where}: items must be pairs")
            if kind == "table":
                cols = block.get("columns")
                if not isinstance(cols, list) or not all(isinstance(r, list) and len(r) == len(cols) for r in block.get("rows", [])):
                    problems.append(f"{where}: rows must match columns")
            if kind == "figure":
                lines = block.get("lines")
                if not isinstance(lines, list) or len(lines) > MAX_LINES:
                    problems.append(f"{where}: at most {MAX_LINES} lines")
                else:
                    for line in lines:
                        x, y = line.get("x"), line.get("y")
                        if not isinstance(x, list) or not isinstance(y, list) or len(x) != len(y):
                            problems.append(f"{where}: x and y must be lists of equal length")
    if not problems and len(json.dumps(doc)) > MAX_BYTES:
        problems.append(f"document larger than {MAX_BYTES} bytes")
    return problems


def figure_png(figure: dict) -> bytes:
    """The figure at 1600x700 px, one line per entry, with a legend and axis labels."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(10.67, 4.67), dpi=150)
    for i, line in enumerate(figure.get("lines", [])):
        ax.plot(line["x"], line["y"], label=line.get("label", ""), color=COLORS[i % len(COLORS)], linewidth=1.8)
    ax.set_xlabel(figure.get("xLabel", ""))
    ax.set_ylabel(figure.get("yLabel", ""))
    ax.grid(True, color="#ddd6cc", linewidth=0.6)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    if figure.get("lines"):
        ax.legend(frameon=False)
    fig.tight_layout()
    out = io.BytesIO()
    fig.savefig(out, format="png")
    plt.close(fig)
    return out.getvalue()


def _table_rows(block: dict) -> tuple[list[str], list[list[str]], str | None]:
    if block["kind"] == "table":
        return block["columns"], block["rows"], block.get("caption")
    return ["Setting", "Value"], block["items"], None


def render_docx(doc: dict) -> bytes:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_heading(doc["title"], level=1)
    document.add_paragraph().add_run(doc["subtitle"]).italic = True
    for section in doc["sections"]:
        document.add_heading(section["heading"], level=2)
        for block in section["blocks"]:
            kind = block["kind"]
            if kind == "paragraph":
                document.add_paragraph(block["text"])
            elif kind == "note":
                document.add_paragraph().add_run(block["text"]).italic = True
            elif kind == "bullets":
                for item in block["items"]:
                    document.add_paragraph(item, style="List Bullet")
            elif kind in ("table", "keyValues"):
                columns, rows, caption = _table_rows(block)
                if caption:
                    document.add_paragraph().add_run(caption).bold = True
                table = document.add_table(rows=1, cols=len(columns))
                table.style = "Table Grid"
                for cell, text in zip(table.rows[0].cells, columns):
                    cell.text = str(text)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                for row in rows:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = str(text)
            elif kind == "figure":
                document.add_picture(io.BytesIO(figure_png(block)), width=Inches(6.3))
                document.add_paragraph().add_run(block["caption"]).italic = True
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_pdf(doc: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    cjk = _has_cjk(json.dumps(doc, ensure_ascii=False))
    font = _rl_register_cjk() if cjk else "Helvetica"
    bold = font if cjk else "Helvetica-Bold"

    def fix(value: Any) -> str:
        text = str(value)
        return text if cjk else _sanitize_latin(text)

    def text(value: Any) -> str:
        return _escape(fix(value))

    styles = {
        "title": ParagraphStyle("title", fontName=bold, fontSize=18, leading=22, spaceAfter=4),
        "subtitle": ParagraphStyle("subtitle", fontName=font, fontSize=10, leading=13, textColor=colors.HexColor("#8a8078"), spaceAfter=14),
        "h2": ParagraphStyle("h2", fontName=bold, fontSize=13, leading=16, spaceBefore=12, spaceAfter=6),
        "body": ParagraphStyle("body", fontName=font, fontSize=10, leading=14, spaceAfter=6),
        "note": ParagraphStyle("note", fontName=font, fontSize=9, leading=12, textColor=colors.HexColor("#8a8078"), spaceAfter=6),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=8.5, leading=11),
        "head": ParagraphStyle("head", fontName=bold, fontSize=8.5, leading=11),
    }
    width = A4[0] - 36 * mm

    def table(columns: list[str], rows: list[list[str]], caption: str | None) -> list:
        flow: list = []
        if caption:
            flow.append(Paragraph(text(caption), styles["note"]))
        data = [[Paragraph(text(c), styles["head"]) for c in columns]]
        data += [[Paragraph(text(c), styles["cell"]) for c in row] for row in rows]
        t = Table(data, colWidths=[width / len(columns)] * len(columns), repeatRows=1)
        t.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#ddd6cc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
        ]))
        flow.append(t)
        flow.append(Spacer(1, 6))
        return flow

    story: list = [Paragraph(text(doc["title"]), styles["title"]), Paragraph(text(doc["subtitle"]), styles["subtitle"])]
    for section in doc["sections"]:
        story.append(Paragraph(text(section["heading"]), styles["h2"]))
        for block in section["blocks"]:
            kind = block["kind"]
            if kind == "paragraph":
                story.append(Paragraph(text(block["text"]), styles["body"]))
            elif kind == "note":
                story.append(Paragraph(text(block["text"]), styles["note"]))
            elif kind == "bullets":
                items = [ListItem(Paragraph(text(i), styles["body"])) for i in block["items"]]
                story.append(ListFlowable(items, bulletType="bullet", bulletFontName=font))
            elif kind in ("table", "keyValues"):
                columns, rows, caption = _table_rows(block)
                story.extend(table(columns, rows, caption))
            elif kind == "figure":
                image = Image(io.BytesIO(figure_png(block)), width=width, height=width * 0.4375)
                story.append(KeepTogether([image, Paragraph(text(block["caption"]), styles["note"])]))
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                      title=fix(doc["title"])).build(story)
    return out.getvalue()
