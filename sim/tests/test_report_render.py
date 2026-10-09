"""The report renderers: a sample document to Word and PDF, read back; the route's contract."""
import importlib
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

import report_render as render

AUTH = {"Authorization": "Bearer test"}

SAMPLE = {
    "version": 1,
    "title": "Measles in Kenya, SIR model, one year",
    "subtitle": "EpiChat report · version 2 · 9 October 2026",
    "generatedAt": "2026-10-09T15:00:00Z",
    "language": "en",
    "sections": [
        {"id": "summary", "heading": "Summary", "blocks": [
            {"kind": "paragraph", "text": "Measles would spread fast; R₀ is high."},
            {"kind": "bullets", "items": ["one", "two"]},
        ]},
        {"id": "results", "heading": "Key results", "blocks": [
            {"kind": "table", "caption": "Results by run", "columns": ["Run", "Peak day"], "rows": [["Run 1", "Day 40"], ["Run 2", "Day 55"]]},
            {"kind": "figure", "id": "infected", "caption": "People infected over time", "xLabel": "Day", "yLabel": "People infected",
             "lines": [{"label": "Run 1", "x": [0, 1, 2, 3], "y": [0, 10, 5, 1]}, {"label": "Run 2", "x": [0, 1, 2, 3], "y": [0, 4, 2, 1]}]},
            {"kind": "note", "text": "Run 3's curve is not available."},
        ]},
        {"id": "modelled", "heading": "What was modelled", "blocks": [{"kind": "keyValues", "items": [["Country", "Kenya"], ["Agents", "10,000"]]}]},
    ],
}


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("SIM_SHARED_SECRET", "test")
    import main as module
    importlib.reload(module)
    return TestClient(module.app)


def test_validate_document_accepts_the_sample_and_names_problems():
    assert render.validate_document(SAMPLE) == []
    bad_figure = {**SAMPLE, "sections": [{"id": "x", "heading": "X", "blocks": [
        {"kind": "figure", "id": "f", "caption": "c", "xLabel": "x", "yLabel": "y", "lines": [{"label": "L", "x": [0, 1], "y": [0]}]}]}]}
    assert any("x and y" in p for p in render.validate_document(bad_figure))
    assert render.validate_document({"version": 2}) != []
    assert render.validate_document("nope") == ["document must be an object"]
    unknown = {**SAMPLE, "sections": [{"id": "s", "heading": "S", "blocks": [{"kind": "video"}]}]}
    assert any("kind" in p for p in render.validate_document(unknown))
    ragged = {**SAMPLE, "sections": [{"id": "s", "heading": "S", "blocks": [{"kind": "table", "columns": ["A", "B"], "rows": [["only one"]]}]}]}
    assert any("rows must match columns" in p for p in render.validate_document(ragged))


def test_docx_carries_headings_tables_and_a_picture():
    data = render.render_docx(SAMPLE)
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == SAMPLE["title"]
    assert "Key results" in texts and "What was modelled" in texts
    assert any(t.rows[0].cells[0].text == "Run" for t in doc.tables)
    assert any(t.rows[1].cells[0].text == "Country" for t in doc.tables)
    assert len(doc.inline_shapes) == 1
    assert "one" in texts and "Run 3's curve is not available." in texts


def test_pdf_renders_latin_and_cjk_text_and_paginates():
    data = render.render_pdf(SAMPLE)
    assert data.startswith(b"%PDF")
    cjk = {**SAMPLE, "sections": [{"id": "summary", "heading": "摘要", "blocks": [{"kind": "paragraph", "text": "麻疹会迅速传播。R₀ 很高。"}]}]}
    assert render.render_pdf(cjk).startswith(b"%PDF")
    long = {**SAMPLE, "sections": [{"id": "summary", "heading": "Summary", "blocks": [{"kind": "paragraph", "text": "word " * 400}] * 12}]}
    assert render.render_pdf(long).count(b"/Type /Page") >= 3


def test_figure_png_is_a_png():
    png = render.figure_png(SAMPLE["sections"][1]["blocks"][1])
    assert png.startswith(b"\x89PNG")


def test_export_route_contract(client):
    assert client.post("/export", json={"format": "pdf", "document": SAMPLE}).status_code == 401
    r = client.post("/export", json={"format": "pdf", "document": SAMPLE}, headers=AUTH)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content.startswith(b"%PDF")
    r = client.post("/export", json={"format": "docx", "document": SAMPLE}, headers=AUTH)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    r = client.post("/export", json={"format": "docx", "document": {"version": 2}}, headers=AUTH)
    assert r.status_code == 422 and r.json()["error"]["kind"] == "invalid_document"
    assert client.post("/export", json={"format": "txt", "document": SAMPLE}, headers=AUTH).status_code == 422


def test_pdf_keeps_non_latin1_text_in_a_unicode_font():
    """A Polish or Vietnamese narrative must not become question marks; R₀ still reads R0 on both paths."""
    assert render.pdf_text("R₀ Łódź Việt Nam Ελλάδα Москва", cjk=False) == "R0 Łódź Việt Nam Ελλάδα Москва"
    assert render.pdf_text("R₀ 麻疹", cjk=True) == "R0 麻疹"
    polish = {**SAMPLE, "sections": [{"id": "summary", "heading": "Podsumowanie", "blocks": [{"kind": "paragraph", "text": "Odra rozprzestrzeniałaby się szybko w Łodzi."}]}]}
    data = render.render_pdf(polish)
    assert data.startswith(b"%PDF")
    assert b"DejaVuSans" in data
    assert b"DejaVuSans" in render.render_pdf(SAMPLE)
